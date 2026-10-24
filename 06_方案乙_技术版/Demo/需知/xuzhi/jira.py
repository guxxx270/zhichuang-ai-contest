"""Jira 工单解析：粘贴的单条工单文本，或 Jira 导出的 CSV / Excel（批量）。

只取评估要用的字段：单号、摘要、类型、状态、优先级、模块、标签、描述（含验收标准 / 技术说明 / 复现与期望段）、
故事点、原始预估、剩余预估、实际耗时。时间统一折成小时：Jira 导出里是秒，粘贴文本里是 3d 4h / 16h / 2天 这类写法。
经办人、报告人等人名字段一律不读（粘贴文本里的这类字段整段丢掉）。"""
from __future__ import annotations

import csv
import io
import re
from dataclasses import dataclass, field
from pathlib import Path

HOURS_PER_DAY = 8.0
DAYS_PER_WEEK = 5.0
# 显式白名单、精确匹配：子串匹配会把「待发布」「未解决」也当成已完成
DONE_STATUSES = ("完成", "已完成", "关闭", "已关闭", "已解决", "已上线", "上线完成", "已发布", "发布完成", "已交付",
                 "验收通过", "已验收", "done", "closed", "resolved")


@dataclass
class JiraIssue:
    key: str = ""
    summary: str = ""
    issue_type: str = ""
    status: str = ""
    priority: str = ""
    components: list[str] = field(default_factory=list)
    labels: list[str] = field(default_factory=list)
    description: str = ""
    acceptance: list[str] = field(default_factory=list)
    tech_notes: str = ""
    repro: str = ""             # 缺陷：复现步骤
    expected: str = ""          # 缺陷：期望结果
    story_points: float | None = None
    original_h: float | None = None
    remaining_h: float | None = None
    spent_h: float | None = None

    @property
    def status_done(self) -> bool:
        """Jira 状态已完成（不管登没登记工时）：批量评估里不算待评估。"""
        return (self.status or "").strip().lower() in DONE_STATUSES

    @property
    def done(self) -> bool:
        """已完成且登记了耗费时间：能对账、能当校准样本。"""
        return self.status_done and bool(self.spent_h)

    @property
    def spent_days(self) -> float | None:
        return round(self.spent_h / HOURS_PER_DAY, 2) if self.spent_h else None

    @property
    def text(self) -> str:
        """规则匹配用的全文：摘要 + 描述 + 验收 + 技术说明 + 复现 / 期望 + 模块。"""
        return "\n".join(x for x in [self.summary, self.description, *self.acceptance, self.tech_notes,
                                     self.repro, self.expected, " ".join(self.components)] if x)

    @property
    def title(self) -> str:
        return f"{self.key} {self.summary}".strip()


# ---------------------------------------------------------------- 时间
# 不加单独的「日」：「10月5日」会被当成 5 天
_DUR = re.compile(r"(\d+(?:\.\d+)?)\s*(人天|人日|工作日|w|周|d|天|h|小时|m(?!s)|分钟?)", re.I)


def parse_duration_h(v) -> float | None:
    """「3d 4h」「16h」「2天」「1.5 人天」「2人日」→ 小时；纯数字按小时。解析不了返回 None。
    括号里是换算或备注（「2d（16h）」），不再重复累加。"""
    s = re.sub(r"[（(][^）)]*[）)]", " ", str(v if v is not None else "")).strip()
    if not s or s.lower() in ("nan", "none", "-", "—"):
        return None
    if re.fullmatch(r"\d+(?:\.\d+)?", s):
        return float(s)
    total, hit = 0.0, False
    for num, unit in _DUR.findall(s):
        n, u = float(num), unit.lower()
        if u in ("w", "周"):
            total += n * DAYS_PER_WEEK * HOURS_PER_DAY
        elif u in ("d", "天", "人天", "人日", "工作日"):
            total += n * HOURS_PER_DAY
        elif u in ("h", "小时"):
            total += n
        else:
            total += n / 60
        hit = True
    return round(total, 2) if hit else None


def _num(v) -> float | None:
    try:
        f = float(str(v).strip())
        return f if f == f else None   # 排除 NaN
    except (TypeError, ValueError):
        return None


# ---------------------------------------------------------------- 文本清洗
def clean_markup(s: str) -> str:
    """Jira wiki 标记 → 纯文本：h3. 标题、{code}/{noformat}、*粗体*、[文字|链接]、{{等宽}}。"""
    s = (s or "").replace("\r\n", "\n").replace("\r", "\n")
    s = re.sub(r"\{(?:code|noformat|quote|panel)(?::[^}]*)?\}", "", s)
    s = re.sub(r"^\s*h[1-6]\.\s*", "", s, flags=re.M)
    s = re.sub(r"\[([^\]|]+)\|[^\]]+\]", r"\1", s)
    s = re.sub(r"\{\{([^}]+)\}\}", r"\1", s)
    s = re.sub(r"(?<![\w*])\*([^*\n]+)\*(?![\w*])", r"\1", s)
    s = s.replace("\\\\", "\n")
    return s.strip()


_BULLET = re.compile(r"^\s*(?:[#*\-•·]+|\d+\s*[.、)）])\s*")


def _items(block: str) -> list[str]:
    out = []
    for ln in (block or "").splitlines():
        t = _BULLET.sub("", ln).strip(" ；;。")
        if t:
            out.append(t)
    return out


# ---------------------------------------------------------------- 粘贴的单条工单
_INLINE = {
    "summary": r"摘要|标题|Summary",
    "issue_type": r"问题类型|类型|Issue\s*Type|Type",
    "status": r"状态|Status",
    "priority": r"优先级|Priority",
    "components": r"模块|组件|Components?|Component/s",
    "labels": r"标签|Labels?",
    "story_points": r"Story\s*Points?|故事点|估算点|SP",
    "original_h": r"原始预估|初始预估|Original\s*Estimate|预估工时|预估",
    "remaining_h": r"剩余预估|Remaining\s*Estimate",
    "spent_h": r"实际耗时|耗费时间|已用时间|已登记工时|实际工时|Time\s*Spent",
}
_SECTIONS = {
    "description": r"描述|Description|背景|需求描述",
    "acceptance": r"验收标准|验收条件|Acceptance\s*Criteria|完成标准|AC",
    "tech_notes": r"技术说明|技术方案|实现说明|技术备注",
    "repro": r"复现步骤|复现|Steps\s*to\s*Reproduce",
    "expected": r"期望结果|期望|Expected(?:\s*Result)?",
}
_PEOPLE = {"_drop": r"经办人|报告人|创建者|创建人|负责人|验收人|处理人|Assignee|Reporter|Creator"}   # 人名字段：认出来就丢掉
_KEY_RE = re.compile(r"\b([A-Z][A-Z0-9]{1,9}-\d{1,6})\b")
# 「UTF-8」「SHA-256」长得像单号，其实是技术词
_NOT_KEY = {"UTF", "SHA", "AES", "RSA", "ISO", "GB", "RFC", "TLS", "SSL", "HTTP", "MD", "CRC", "UTC", "GMT"}
# 段落标题后面要么是冒号，要么整行就是标题——「Active users…」不能被当成 AC 段；
# 标题前后可带 Markdown 的 #、【】、*粗体* 和「（AC）」这类括注
_SEC_RE = re.compile(r"^\s*(?:h[1-6]\.|#{1,6})?\s*[*【\[]*\s*(" + "|".join(_SECTIONS.values())
                     + r")\s*(?:[（(][^）)]{0,12}[）)])?\s*[*】\]]*\s*(?:[:：]\s*(.*))?$", re.I)
_LABEL_ONLY = re.compile(r"\s*([A-Za-z一-鿿 /]{1,24}?)\s*[:：]\s*")   # 「类型:」单独一行，值在下一行
# 没有验收段时，描述里的 Given / When / Then 才算验收；「当前…」「当日…」「假设用户…」不算
_GWT = re.compile(r"(?i)(?:given|when|then)\b|当(?!前|日|天|月|年).+时|假设.+[，,]")


def _which(patterns: dict[str, str], label: str) -> str | None:
    for k, p in patterns.items():
        if re.fullmatch(p, label.strip(), flags=re.I):
            return k
    return None


def _field_label(label: str) -> str | None:
    return _which(_INLINE, label) or _which(_PEOPLE, label)


def _inline_fields(line: str) -> dict[str, str]:
    """一行里「类型：Story　优先级：High　模块：报表平台」这类字段；以全角空格、竖线、Tab 或两个以上空格分隔。"""
    line = re.sub(r"([:：])[ \t]*\t[ \t]*", r"\1 ", line)   # 「类型:<Tab>故事」：冒号后的 Tab 是标签和值之间的，不是字段之间的
    found: dict[str, str] = {}
    for seg in re.split(r"[　|｜\t]|\s{2,}", line):
        m = re.match(r"^\s*([A-Za-z一-鿿 /]{1,24}?)\s*[:：]\s*(.+?)\s*$", seg)
        if m:
            k = _field_label(m.group(1))
            if k:
                found[k] = m.group(2)
    return found


def _section_of(line: str) -> str | None:
    sm = _SEC_RE.match(line)
    return _which(_SECTIONS, sm.group(1)) if sm else None


def _apply_inline(issue: JiraIssue, k: str, v: str) -> None:
    if k == "_drop":
        return
    if k in ("original_h", "remaining_h", "spent_h"):
        setattr(issue, k, parse_duration_h(v))
    elif k == "story_points":   # 只取第一个数：「5（约 2 天）」是 5 点，不是 52
        m = re.search(r"\d+(?:\.\d+)?", v)
        issue.story_points = float(m.group()) if m else None
    elif k in ("components", "labels"):
        setattr(issue, k, [x.strip() for x in re.split(r"[/、,，;；]", v) if x.strip()])
    else:
        setattr(issue, k, v.strip())


def _find_key(lines: list[str]) -> str:
    """单号只在前 3 个非空行里找（详情页的面包屑「项目 / XZ-1050」也算），正文里引用的「复用 XZ-1005」不算。"""
    for ln in lines[:3]:
        for m in _KEY_RE.finditer(ln):
            if m.group(1).split("-")[0] not in _NOT_KEY:
                return m.group(1)
    return ""


def _strip_key(line: str, key: str) -> str:
    """去掉单号和「【JIRA】」后剩下的标题文字；只删选中的那个单号，「SHA-256」这类保留。
    面包屑「项目 / XZ-1050」去掉单号后只剩项目路径，不当摘要。"""
    rest = line.replace(key, "") if key else line
    if key and key in line and re.search(r"[/>›]\s*$", rest):
        return ""
    return re.sub(r"【?\s*JIRA\s*】?", "", rest, flags=re.I).strip(" :：-—[]【】")


def parse_issue_text(text: str, inline: bool = True) -> JiraIssue:
    """从 Jira 详情页复制来的文字解析一条工单（字段缺什么就空着，评估时按缺失处理）。
    inline=False 只认段落标题：导出里的描述正文交给这里拆段时，「状态：线上已复现」这类行要原样留在描述里。"""
    raw = clean_markup(text)
    issue = JiraIssue()
    lines = [ln for ln in raw.splitlines() if ln.strip()]
    issue.key = _find_key(lines)
    sections: dict[str, list[str]] = {}
    current = None
    label = None   # 上一行是孤立的字段标签（「类型:」），这一行是它的值
    head_lines: list[str] = []
    for ln in lines:
        if label:
            k, label = label, None
            if not _inline_fields(ln) and not _section_of(ln):
                _apply_inline(issue, k, ln)
                continue
        if inline:
            m = _LABEL_ONLY.fullmatch(ln)
            if m and _field_label(m.group(1)):
                label, current = _field_label(m.group(1)), None
                continue
            fields = _inline_fields(ln)
            if fields:
                for k, v in fields.items():
                    _apply_inline(issue, k, v)
                current = None
                continue
        sm = _SEC_RE.match(ln)
        if sm and _which(_SECTIONS, sm.group(1)):
            current = _which(_SECTIONS, sm.group(1))
            if (sm.group(2) or "").strip():
                sections.setdefault(current, []).append(sm.group(2).strip())
            continue
        if current:
            sections.setdefault(current, []).append(ln)
        else:
            head_lines.append(ln)
    titles = [_strip_key(ln, issue.key) for ln in head_lines]
    summary_idx = None
    if not issue.summary:
        summary_idx = next((i for i, t in enumerate(titles) if t), None)
        issue.summary = titles[summary_idx] if summary_idx is not None else ""
    # 只排除用作摘要的那一行（以及只有单号的行）；摘要来自「摘要：」字段时，正文第一行照样留在描述里
    rest = [ln for i, ln in enumerate(head_lines) if i != summary_idx and titles[i] and ln.strip() != issue.summary]
    issue.description = "\n".join(sections.get("description", []) + rest).strip()
    issue.acceptance = _items("\n".join(sections.get("acceptance", [])))
    if not issue.acceptance:   # 没有验收段：描述里 Given / When / Then 也算
        issue.acceptance = [x for x in _items(issue.description) if _GWT.match(x)]
    issue.tech_notes = "\n".join(sections.get("tech_notes", [])).strip()
    issue.repro = "\n".join(sections.get("repro", [])).strip()
    issue.expected = "\n".join(sections.get("expected", [])).strip()
    return issue


def _split_description(issue: JiraIssue, desc: str) -> None:
    """导出里的描述是整段 wiki 文本：按段落标题拆出验收标准、技术说明、复现与期望。"""
    parsed = parse_issue_text("描述：\n" + (desc or ""), inline=False)
    issue.description = parsed.description
    issue.acceptance = issue.acceptance or parsed.acceptance
    issue.tech_notes, issue.repro, issue.expected = parsed.tech_notes, parsed.repro, parsed.expected


# ---------------------------------------------------------------- 导出的 CSV / Excel
_COLS = {
    "key": ("问题关键字", "关键字", "issue key", "key", "单号"),
    "summary": ("摘要", "summary", "标题"),
    "issue_type": ("问题类型", "issue type", "类型"),
    "status": ("状态", "status"),
    "priority": ("优先级", "priority"),
    "components": ("模块", "组件", "component/s", "components", "component"),
    "labels": ("标签", "labels", "label"),
    "description": ("描述", "description"),
    "original_h": ("原始预估", "original estimate", "初始预估"),
    "remaining_h": ("剩余预估", "remaining estimate"),
    "spent_h": ("耗费时间", "已用时间", "实际耗时", "time spent", "已登记工时"),
}


def _norm_header(h: str) -> str:
    h = (h or "").strip().lower()
    h = re.sub(r"^(custom field|自定义字段)\s*\((.*)\)$", r"\2", h)
    return h.strip()


_ACCEPT_COLS = ("验收标准", "验收条件", "acceptance criteria")   # 精确匹配：「验收人」「验收时间」不是验收标准


def _field_of(header: str) -> str | None:
    h = _norm_header(header)
    if h.startswith("σ") or h.startswith("Σ".lower()):
        return None   # Σ 汇总列（含子任务）不要
    for k, names in _COLS.items():
        if h in names:
            return k
    if re.search(r"父|parent|子任务|sub-?task|史诗|epic", h):
        return None   # 「父级摘要」「子任务预估」说的不是本条
    if "story point" in h or "故事点" in h:
        return "story_points"
    if h in _ACCEPT_COLS:
        return "acceptance"
    # 自己整理的表、不同语言的 Jira：按关键词归类（「预估工时（小时）」「Time Spent (hours)」「概要」）
    if re.search(r"剩余|remaining", h):
        return "remaining_h" if re.search(r"预估|估算|estimate", h) else None
    if re.search(r"原始预估|初始预估|预估工时|original estimate", h):
        return "original_h"
    if re.search(r"耗费时间|实际工时|实际耗时|已用时间|已登记工时|time spent", h):
        return "spent_h"
    if re.search(r"摘要|概要|主题|标题|summary", h):
        return "summary"
    return None


def _hours(header: str, v) -> float | None:
    if v is None or str(v).strip() == "":
        return None
    n = _num(v)
    if n is None:
        return parse_duration_h(v)
    if re.search(r"人天|人日|days?\b|\(d\)|（d）", header, flags=re.I):
        return round(n * HOURS_PER_DAY, 2)
    if re.search(r"小时|\(h\)|（h）|hours?", header, flags=re.I):
        return round(n, 2)
    return round(n / 3600, 2)   # Jira 导出的时间字段是秒


def _rows_from_bytes(data: bytes, name: str = "") -> list[list[str]]:
    if name.lower().endswith((".xlsx", ".xlsm")):
        from openpyxl import load_workbook

        ws = load_workbook(io.BytesIO(data), read_only=True, data_only=True).active
        return [["" if c is None else str(c) for c in row] for row in ws.iter_rows(values_only=True)]
    for enc in ("utf-8-sig", "gb18030"):
        try:
            return list(csv.reader(io.StringIO(data.decode(enc))))
        except UnicodeDecodeError:
            continue
    return list(csv.reader(io.StringIO(data.decode("utf-8", errors="replace"))))


def parse_export(data: bytes, name: str = "") -> list[JiraIssue]:
    """Jira 筛选结果导出的 CSV（同名列可重复，如多个「模块」）或另存的 Excel → 工单列表。"""
    rows = [r for r in _rows_from_bytes(data, name) if any(str(c).strip() for c in r)]
    if not rows:
        return []
    header = rows[0]
    fields = [_field_of(h) for h in header]
    issues = []
    for row in rows[1:]:
        it = JiraIssue()
        desc = ""
        for h, k, v in zip(header, fields, row):
            v = (v or "").strip()
            if not k or not v:
                continue
            if k in ("components", "labels"):
                getattr(it, k).extend(x.strip() for x in re.split(r"[,，;；]", v) if x.strip())
            elif k == "description":
                desc = v
            elif k == "acceptance":
                it.acceptance = _items(clean_markup(v))
            elif k in ("original_h", "remaining_h", "spent_h"):
                setattr(it, k, _hours(h, v))
            elif k == "story_points":
                it.story_points = _num(v)
            else:
                setattr(it, k, v)
        _split_description(it, desc)
        if it.summary or it.key:
            issues.append(it)
    return issues


def load_export(path: str | Path) -> list[JiraIssue]:
    p = Path(path)
    return parse_export(p.read_bytes(), p.name)
